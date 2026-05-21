from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from django.db.models import Count, Q
from django.utils import timezone
from datetime import datetime
from accounts.models import User
from accounts.models import Booking
from core.models import CategoryService


def export_dashboard_excel(request):
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    ws1 = wb.create_sheet(title="Users Analytics")

    headers = ["Customers", "Vendors", "Admins"]

    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)

    user_counts = User.objects.aggregate(
        total_customers=Count("id", filter=Q(role="customer")),
        total_vendors=Count("id", filter=Q(role="vendor")),
        total_admins=Count("id", filter=Q(role="admin")),
    )

    ws1.append([
        user_counts["total_customers"],
        user_counts["total_vendors"],
        user_counts["total_admins"],
    ])



    ws2 = wb.create_sheet(title="Booking Analytics")

    booking_headers = [
        "Total",
        "Pending",
        "Assigned",
        "Accepted",
        "In Progress",
        "Completed",
        "Cancelled",
    ]

    for col_num, header in enumerate(booking_headers, 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)

    bookings = Booking.objects.all()

    booking_counts = bookings.aggregate(
        total_bookings=Count("id"),
        total_pending=Count("id", filter=Q(status="pending")),
        total_assigned=Count("id", filter=Q(status="assigned")),
        total_accepted=Count("id", filter=Q(status="accepted")),
        total_in_progress=Count("id", filter=Q(status="in_progress")),
        total_completed=Count("id", filter=Q(status="completed")),
        total_cancelled=Count("id", filter=Q(status="cancelled")),
    )

    ws2.append([
        booking_counts["total_bookings"],
        booking_counts["total_pending"],
        booking_counts["total_assigned"],
        booking_counts["total_accepted"],
        booking_counts["total_in_progress"],
        booking_counts["total_completed"],
        booking_counts["total_cancelled"],
    ])



    ws3 = wb.create_sheet(title="Monthly Leads")

    ws3["A1"] = "Month"
    ws3["B1"] = "Bookings"

    ws3["A1"].font = Font(bold=True)
    ws3["B1"].font = Font(bold=True)

    current_year = timezone.now().year

    for month in range(1, 13):

        total = bookings.filter(
            booking_created_at__year=current_year,
            booking_created_at__month=month
        ).count()

        month_name = datetime(current_year, month, 1).strftime("%b")

        ws3.append([month_name, total])

   

    ws4 = wb.create_sheet(title="Service Analytics")

    ws4["A1"] = "Service"
    ws4["B1"] = "Bookings"

    ws4["A1"].font = Font(bold=True)
    ws4["B1"].font = Font(bold=True)

    services = CategoryService.objects.all()

    for service in services:

        total = bookings.filter(service=service).count()

        ws4.append([
            service.s_title,
            total
        ])



    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = (
        'attachment; filename=dashboard_analytics.xlsx'
    )

    wb.save(response)

    return response

